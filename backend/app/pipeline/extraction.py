import fitz # PyMuPDF
import docx
import openpyxl
from typing import Tuple, Dict, Any

def extract_text_from_file(file_path: str, file_info: Dict[str, Any]) -> Tuple[str, bool, Dict[str, Any]]:
    """
    Extracts text directly from digital formats (PDF, DOCX, XLSX, TXT).
    Returns: (extracted_text, is_ocr_needed, metadata)
    """
    ext = file_info["extension"]
    extracted_text = ""
    is_ocr_needed = False
    metadata = {}

    try:
        if ext == "txt":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                extracted_text = f.read()
            is_ocr_needed = False
            metadata["page_count"] = 1

        elif ext == "docx":
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Also extract tables
            table_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        table_text.append(" | ".join(row_data))

            extracted_text = "\n".join(paragraphs + table_text)
            is_ocr_needed = False
            metadata["page_count"] = len(doc.paragraphs) // 10 + 1

        elif ext == "xlsx":
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet_contents = []
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                sheet_text = [f"--- Sheet: {sheet_name} ---"]
                for row in sheet.iter_rows(values_only=True):
                    row_str = " | ".join([str(val) for val in row if val is not None])
                    if row_str.strip():
                        sheet_text.append(row_str)
                sheet_contents.append("\n".join(sheet_text))
            
            extracted_text = "\n\n".join(sheet_contents)
            is_ocr_needed = False
            metadata["sheet_names"] = wb.sheetnames

        elif ext == "pdf":
            doc = fitz.open(file_path)
            metadata["page_count"] = len(doc)
            pages_text = []
            
            for page in doc:
                text = page.get_text("text")
                if text.strip():
                    pages_text.append(text.strip())

            extracted_text = "\n\n".join(pages_text)

            # Detect if digital PDF text is sufficient or if OCR is required
            cleaned_check = "".join(extracted_text.split())
            if len(cleaned_check) < 30: # Scanned PDF or image-only PDF
                is_ocr_needed = True
                extracted_text = ""
            else:
                is_ocr_needed = False

        elif ext in ["png", "jpg", "jpeg"]:
            is_ocr_needed = True
            extracted_text = ""
            metadata["page_count"] = 1

    except Exception as e:
        print(f"[Extraction Error] Error extracting from {file_path}: {e}")
        is_ocr_needed = True

    return extracted_text.strip(), is_ocr_needed, metadata
